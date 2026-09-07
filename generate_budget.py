import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


def create_budget_document():
    doc = docx.Document()

    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Document Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run(
        "MDMIS – 2-Week Backend Development Budget & Technical Plan"
    )
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)

    # Subtitle / Overview
    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run(
        "Stack: Python FastAPI + Supabase (PostgreSQL / PostGIS)\nDuration: 2-Week MVP Sprint | Team Size: 3 Backend Developers"
    )
    r_meta.font.name = "Arial"
    r_meta.font.size = Pt(10)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph()  # Spacer

    # Section Header
    h1 = doc.add_paragraph()
    r_h1 = h1.add_run("1. Budget Allocation Breakdown")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(0, 51, 102)

    # Table Data Definition
    table_data = [
        [
            "Category / Item",
            "Description & Tooling",
            "Est. Cost (USD)",
            "Notes / Allocation",
        ],
        [
            "Developer Stipends",
            "3 Backend Developers (2-Week Sprint)",
            "$300.00 – $600.00",
            "Option for $100–$200/dev based on project tier",
        ],
        [
            "AI Assistant",
            "Claude Pro Plan (Shared Account)",
            "$20.00",
            "FastAPI boilerplate, PostGIS queries, and async code",
        ],
        [
            "Database & Auth",
            "Supabase (PostgreSQL + PostGIS)",
            "$0.00",
            "Free Tier (500 MB DB, 1 GB Storage, Auth included)",
        ],
        [
            "3D Map & Spatial Tiles",
            "3D Mapper / Mapbox API",
            "$0.00 – $15.00",
            "Free Tier dev usage allowance for initial testing",
        ],
        [
            "Object Storage",
            "Cloudflare R2 / AWS S3",
            "$0.00 – $5.00",
            "Host raw GeoTIFF, OBJ, and 3D point cloud assets",
        ],
        [
            "FastAPI Server Hosting",
            "Render / DigitalOcean VPS",
            "$0.00 – $7.00",
            "Free Tier on Render or $6/mo Docker Linux Droplet",
        ],
        [
            "TOTAL ESTIMATED SPEND",
            "Combined Infrastructure + Team",
            "$320.00 – $647.00",
            "Lean Launch Range",
        ],
    ]

    # Create Table
    table = doc.add_table(rows=len(table_data), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Populate and Format Table
    for row_idx, row in enumerate(table_data):
        for col_idx, cell_value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = cell_value
            p = cell.paragraphs[0]
            r = p.runs[0]
            r.font.name = "Arial"
            r.font.size = Pt(9.5)

            # Header Formatting
            if row_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                # Background color for header
                shading_elm = docx.oxml.parse_xml(
                    r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="003366"/>'
                )
                cell._tc.get_or_add_tcPr().append(shading_elm)

            # Footer / Total Formatting
            elif row_idx == len(table_data) - 1:
                r.font.bold = True
                shading_elm = docx.oxml.parse_xml(
                    r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="F0F0F0"/>'
                )
                cell._tc.get_or_add_tcPr().append(shading_elm)

    doc.add_paragraph()  # Spacer

    # Section Header 2
    h2 = doc.add_paragraph()
    r_h2 = h2.add_run("2. Technical Dependencies & Recommendations")
    r_h2.font.name = "Arial"
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(0, 51, 102)

    bullets = [
        (
            "Spatial Database Capabilities",
            "Enable PostGIS extension on Supabase (`CREATE EXTENSION postgis;`) for efficient 3D spatial indexing and coordinate queries.",
        ),
        (
            "Asynchronous Python Framework",
            "Use FastAPI paired with `asyncpg`, `SQLAlchemy`, and `geoalchemy2` for non-blocking spatial queries.",
        ),
        (
            "Storage Architecture",
            "Store spatial metadata in PostgreSQL, but stream large 3D glTF/GLB models directly via Cloudflare R2 presigned URLs.",
        ),
        (
            "DevOps & Local Dev",
            "Utilize Docker Compose locally to ensure zero infrastructure cost prior to final deployment.",
        ),
    ]

    for b_title, b_desc in bullets:
        bp = doc.add_paragraph(style="List Bullet")
        r_title = bp.add_run(f"{b_title}: ")
        r_title.bold = True
        r_title.font.name = "Arial"
        r_title.font.size = Pt(10)

        r_desc = bp.add_run(b_desc)
        r_desc.font.name = "Arial"
        r_desc.font.size = Pt(10)

    # Save Document
    filename = "MDMIS_2Week_Backend_Budget.docx"
    doc.save(filename)
    print(f"Document successfully created: {filename}")


if __name__ == "__main__":
    create_budget_document()